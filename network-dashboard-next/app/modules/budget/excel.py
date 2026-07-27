from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_FORMAT = "network-dashboard-budget-xlsx-v1"
_MAX_BYTES = 15 * 1024 * 1024
_MAX_SHEETS = 100
_MAX_CELLS_PER_SHEET = 10_000
_INVALID_TITLE = re.compile(r"[\\/*?:\[\]]")


def _safe_title(value: str, used: set[str]) -> str:
    base = _INVALID_TITLE.sub("-", str(value or "Budget")).strip(" '")[:31] or "Budget"
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        tail = f"-{suffix}"
        candidate = base[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


class BudgetExcelWorkbook:
    @staticmethod
    def export(payload: dict[str, Any]) -> bytes:
        if payload.get("format") != "network-dashboard-budget-v1":
            raise ValueError("Ogiltigt budgetformat")
        workbook = Workbook()
        workbook.remove(workbook.active)
        used: set[str] = set()
        meta = workbook.create_sheet("_LindellsMeta")
        meta.sheet_state = "hidden"
        meta.append(["format", _FORMAT])
        meta.append(["source_format", payload["format"]])
        meta.append(["exported_at", str(payload.get("exported_at") or "")])
        meta.append(["sheet_title", "sheet_name", "kind", "sort_order", "max_row", "max_col"])

        for order, sheet in enumerate(list(payload.get("sheets") or [])[:_MAX_SHEETS]):
            if not isinstance(sheet, dict):
                continue
            original_name = str(sheet.get("name") or f"Budget {order + 1}")[:80]
            title = _safe_title(original_name, used)
            worksheet = workbook.create_sheet(title)
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            meta.append([
                title,
                original_name,
                str(sheet.get("kind") or "month")[:30],
                int(sheet.get("sort_order") or order),
                int(sheet.get("max_row") or 40),
                int(sheet.get("max_col") or 12),
            ])
            cells = list(sheet.get("cells") or [])[:_MAX_CELLS_PER_SHEET]
            for item in cells:
                if not isinstance(item, dict):
                    continue
                row_num = int(item.get("row_num") or 0)
                col_num = int(item.get("col_num") or 0)
                if 1 <= row_num <= 10_000 and 1 <= col_num <= 1_000:
                    worksheet.cell(row=row_num, column=col_num, value=_cell_value(item.get("value")))
            max_col = min(max(int(sheet.get("max_col") or 12), 1), 30)
            for column in range(1, max_col + 1):
                width = 14
                values = [worksheet.cell(row=row, column=column).value for row in range(1, min(80, worksheet.max_row) + 1)]
                text_lengths = [len(str(value)) for value in values if value not in (None, "")]
                if text_lengths:
                    width = max(10, min(36, max(text_lengths) + 2))
                worksheet.column_dimensions[get_column_letter(column)].width = width
            for cell in worksheet[1]:
                if cell.value not in (None, ""):
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="174C39")
                    cell.alignment = Alignment(vertical="center")

        if len(workbook.sheetnames) == 1:
            workbook.create_sheet("Budget")
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    @staticmethod
    def parse(data: bytes) -> dict[str, Any]:
        if not data:
            raise ValueError("Excel-filen är tom")
        if len(data) > _MAX_BYTES:
            raise ValueError("Excel-filen är för stor")
        try:
            workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
        except Exception as exc:
            raise ValueError("Excel-filen kunde inte läsas") from exc
        if "_LindellsMeta" not in workbook.sheetnames:
            raise ValueError("Filen saknar Lindells versionsmetadata")
        meta = workbook["_LindellsMeta"]
        if meta["B1"].value != _FORMAT:
            raise ValueError("Excel-formatet stöds inte")

        mappings: dict[str, dict[str, Any]] = {}
        for row in meta.iter_rows(min_row=5, values_only=True):
            title, name, kind, sort_order, max_row, max_col = (list(row) + [None] * 6)[:6]
            if not title:
                continue
            mappings[str(title)] = {
                "name": str(name or title)[:80],
                "kind": str(kind or "month")[:30],
                "sort_order": int(sort_order or 0),
                "max_row": max(1, min(10_000, int(max_row or 40))),
                "max_col": max(1, min(1_000, int(max_col or 12))),
            }

        sheets: list[dict[str, Any]] = []
        total_cells = 0
        for title in workbook.sheetnames:
            if title == "_LindellsMeta":
                continue
            if len(sheets) >= _MAX_SHEETS:
                raise ValueError("Excel-filen innehåller för många blad")
            worksheet = workbook[title]
            metadata = mappings.get(title, {
                "name": title[:80],
                "kind": "month",
                "sort_order": len(sheets),
                "max_row": min(max(worksheet.max_row, 1), 10_000),
                "max_col": min(max(worksheet.max_column, 1), 1_000),
            })
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value in (None, ""):
                        continue
                    cells.append({
                        "row_num": cell.row,
                        "col_num": cell.column,
                        "value": _cell_value(cell.value)[:4000],
                    })
                    if len(cells) > _MAX_CELLS_PER_SHEET:
                        raise ValueError(f"Bladet {metadata['name']} innehåller för många celler")
            total_cells += len(cells)
            sheets.append({**metadata, "cells": cells})
        return {
            "format": "network-dashboard-budget-v1",
            "exported_at": str(meta["B3"].value or ""),
            "sheets": sheets,
            "summary": {
                "sheets": len(sheets),
                "cells": total_cells,
                "names": [sheet["name"] for sheet in sheets],
            },
        }
