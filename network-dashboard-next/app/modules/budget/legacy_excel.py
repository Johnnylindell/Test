from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import load_workbook

from app.modules.budget.excel import BudgetExcelWorkbook
from app.modules.budget.xlsx_safety import validate_xlsx_archive

_MAX_SHEETS = 100
_MAX_CELLS_PER_SHEET = 10_000
_MONTH_WORDS = {
    "januari", "februari", "mars", "april", "maj", "juni", "juli", "augusti", "september", "oktober", "november", "december",
    "january", "february", "march", "may", "june", "july", "august", "october",
    "tammikuu", "helmikuu", "maaliskuu", "huhtikuu", "toukokuu", "kesäkuu", "kesakuu", "heinäkuu", "heinakuu", "elokuu", "syyskuu", "lokakuu", "marraskuu", "joulukuu",
}


def _value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _kind(title: str) -> str:
    normalized = re.sub(r"[^a-zåäö0-9]+", " ", title.casefold()).strip()
    if any(month in normalized for month in _MONTH_WORDS):
        return "month"
    if re.search(r"\b(19|20)\d{2}\b", normalized) or any(word in normalized for word in ("year", "år", "vuosi", "summary", "sammanfattning")):
        return "year"
    return "month"


class LegacyBudgetWorkbook:
    @staticmethod
    def parse(data: bytes) -> dict[str, Any]:
        validate_xlsx_archive(data)
        try:
            workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
        except Exception as exc:
            raise ValueError("Excel-filen kunde inte läsas") from exc
        if "_LindellsMeta" in workbook.sheetnames:
            raise ValueError("Arbetsboken innehåller Lindells-metadata men formatversionen kunde inte läsas")

        sheets = []
        total_cells = 0
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            if len(sheets) >= _MAX_SHEETS:
                raise ValueError("Excel-filen innehåller för många synliga blad")
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value in (None, ""):
                        continue
                    cells.append({
                        "row_num": int(cell.row),
                        "col_num": int(cell.column),
                        "value": _value(cell.value)[:4000],
                    })
                    if len(cells) > _MAX_CELLS_PER_SHEET:
                        raise ValueError(f"Bladet {worksheet.title} innehåller för många använda celler")
            total_cells += len(cells)
            sheets.append({
                "name": worksheet.title[:80],
                "kind": _kind(worksheet.title),
                "sort_order": len(sheets),
                "max_row": max(1, min(10_000, int(worksheet.max_row or 1))),
                "max_col": max(1, min(1_000, int(worksheet.max_column or 1))),
                "cells": cells,
            })
        if not sheets:
            raise ValueError("Excel-filen innehåller inga synliga budgetblad")
        return {
            "format": "network-dashboard-budget-v1",
            "source_format": "ordinary-xlsx",
            "exported_at": datetime.now(timezone.utc).isoformat(),
            "sheets": sheets,
            "summary": {
                "sheets": len(sheets),
                "cells": total_cells,
                "names": [sheet["name"] for sheet in sheets],
                "ordinary_workbook": True,
            },
        }


def parse_excel_workbook(data: bytes) -> dict[str, Any]:
    validate_xlsx_archive(data)
    try:
        parsed = BudgetExcelWorkbook.parse(data)
        parsed["source_format"] = "lindells-versioned-xlsx"
        parsed.setdefault("summary", {})["ordinary_workbook"] = False
        return parsed
    except ValueError as exc:
        if "versionsmetadata" not in str(exc) and "versionsmetadata" not in str(exc).casefold():
            raise
    return LegacyBudgetWorkbook.parse(data)
